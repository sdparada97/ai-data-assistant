"""Exploración reproducible de las fuentes de Brío Analyst (ETAPA 4).

Perfila el CSV de ventas, detecta inconsistencias con reglas genéricas, mide su
impacto en la venta neta de la marca indicada, calcula la variación entre dos
períodos por región y extrae el texto de los documentos.

Uso (desde la raíz del repositorio):

    uv run --all-packages --group explore python backend/scripts/explore_data.py \\
        --raw-dir data/raw --sales-file 02_Ventas_2026.csv --brand Brio \\
        --period T1=2026-01,2026-02,2026-03 --period T2=2026-04,2026-05,2026-06 \\
        --output-dir data/processed/exploration

No modifica ``data/raw``. Escribe ``report.md`` y el texto de cada documento en
``--output-dir``.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import Path

import docx
import pandas as pd
import pymupdf

CATEGORICAL = ["mes", "marca", "region", "canal", "punto_venta_id", "presentacion"]
MONTH_PATTERN = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
PDV_PATTERN = re.compile(r"^(?P<region>[A-Z]+)-(?P<canal>[A-Z])-\d+$")
PRICE_DEVIATION = 0.5  # desviación relativa frente a la mediana marca/presentación


@dataclass(frozen=True)
class Rule:
    code: str
    name: str
    proposal: str
    detect: Callable[[pd.DataFrame], pd.Series]


# --- carga y utilidades ------------------------------------------------------


def load_raw(path: Path) -> pd.DataFrame:
    """Lee todo como texto para no ocultar inconsistencias al convertir tipos."""
    return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")


def _numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce")


def _canonical_spelling(series: pd.Series) -> dict[str, str]:
    """Mapea cada variante (sin espacios, sin mayúsculas) a su escritura más frecuente."""
    stripped = series.str.strip()
    canonical = stripped.groupby(stripped.str.lower()).agg(lambda s: s.value_counts().idxmax())
    return {value: canonical[value.strip().lower()] for value in series.unique()}


# --- reglas de calidad -------------------------------------------------------


def build_rules(raw: pd.DataFrame) -> list[Rule]:
    # DQ09 y DQ10 (coherencia con el código del punto de venta) se omitieron: el patrón
    # RR-C-NNN se deduce de los datos, pero ningún documento lo define.
    pdv_counts = raw["punto_venta_id"].str.strip().value_counts()

    def unit_price_outlier(df: pd.DataFrame) -> pd.Series:
        units, sales = _numeric(df["unidades"]), _numeric(df["venta_neta_cop"])
        price = (sales / units).where((units > 0) & (sales > 0))
        keys = [df["marca"].str.strip().str.lower(), df["presentacion"].str.strip()]
        median = price.groupby(keys).transform("median")
        return ((price / median - 1).abs() > PRICE_DEVIATION).fillna(False)

    return [
        Rule("DQ01", "Fila duplicada exacta", "Eliminar duplicados exactos y conservar la primera aparición.",
             lambda df: df.duplicated(keep="first")),
        Rule("DQ02", "registro_id repetido con contenido distinto", "Marcar para revisión; no sumar ambas.",
             lambda df: df["registro_id"].duplicated(keep=False) & ~df.duplicated(keep=False)),
        Rule("DQ03", "Espacios sobrantes en dimensiones", "Recortar espacios al inicio y al final.",
             lambda df: pd.concat([df[c] != df[c].str.strip() for c in CATEGORICAL], axis=1).any(axis=1)),
        Rule("DQ04", "Marca con escritura no canónica", "Normalizar a la escritura más frecuente (sin distinguir mayúsculas).",
             lambda df: df["marca"] != df["marca"].map(_canonical_spelling(df["marca"]))),
        Rule("DQ05", "Mes con formato o valor inválido", "Excluir del análisis por período y reportar (no se puede asignar a T1/T2).",
             lambda df: ~df["mes"].str.strip().str.match(MONTH_PATTERN)),
        Rule("DQ06", "Unidades vacías o no numéricas", "Excluir la fila completa y reportar.",
             lambda df: _numeric(df["unidades"]).isna()),
        Rule("DQ07", "Venta neta vacía o no numérica", "Excluir la fila completa, sin imputar, y reportar.",
             lambda df: _numeric(df["venta_neta_cop"]).isna()),
        Rule("DQ08", "Valores negativos", "Excluir y reportar: no hay política para devoluciones o ajustes.",
             lambda df: (_numeric(df["unidades"]) < 0) | (_numeric(df["venta_neta_cop"]) < 0)),
        Rule("DQ11", "Punto de venta con una sola aparición", "Solo informativo: indicio de registro anómalo.",
             lambda df: df["punto_venta_id"].str.strip().map(pdv_counts).eq(1)),
        Rule("DQ12", "Precio unitario atípico (> 50 % de la mediana marca/presentación)", "Solo informativo: revisar el registro.",
             unit_price_outlier),
    ]


def clean(raw: pd.DataFrame, *, normalize: bool = True, dedup: bool = True,
          impute_missing_sales: bool = False) -> pd.DataFrame:
    """Aplica las reglas propuestas. Devuelve columnas numéricas ``u`` y ``v``.

    Los indicadores permiten medir la sensibilidad del resultado a cada regla.
    ``impute_missing_sales`` es solo un escenario de comparación, no una regla propuesta.
    """
    df = raw.copy()
    if normalize:
        for column in CATEGORICAL:
            df[column] = df[column].str.strip()
        df["marca"] = df["marca"].map(_canonical_spelling(df["marca"]))
    if dedup:
        df = df.drop_duplicates(keep="first")
    df = df[df["mes"].str.strip().str.match(MONTH_PATTERN)]
    df["u"], df["v"] = _numeric(df["unidades"]), _numeric(df["venta_neta_cop"])
    if impute_missing_sales:
        price = (df["v"] / df["u"]).where((df["u"] > 0) & (df["v"] > 0))
        keys = [df["marca"], df["region"], df["presentacion"], df["mes"]]
        df["v"] = df["v"].fillna(df["u"] * price.groupby(keys).transform("median"))
    # DQ06/DQ07 (aprobadas): una fila sin unidades o sin venta neta válidas se excluye completa.
    df = df[df["v"].notna() & df["u"].notna()]
    return df[~((df["u"] < 0) | (df["v"] < 0))]


def assign_period(df: pd.DataFrame, periods: dict[str, list[str]]) -> pd.Series:
    month_to_period = {m: name for name, months in periods.items() for m in months}
    return df["mes"].map(month_to_period)


# --- secciones del reporte ---------------------------------------------------


def _md(frame: pd.DataFrame | pd.Series) -> str:
    """Tabla Markdown: importes enteros con separador de miles y porcentajes con 2 decimales."""
    table = frame.to_frame() if isinstance(frame, pd.Series) else frame.copy()
    for column in table.columns:
        if pd.api.types.is_numeric_dtype(table[column]) and "pct" not in str(column):
            if table[column].dropna().abs().max() >= 1000:
                table[column] = table[column].round(0).astype("Int64")
    return table.to_markdown(floatfmt=".2f", intfmt=",")


def _fmt(frame: pd.DataFrame) -> str:
    return "```\n" + frame.to_string() + "\n```"


def profile(raw: pd.DataFrame) -> str:
    lines = [f"Filas: {len(raw)} · Columnas: {', '.join(raw.columns)}", ""]
    for column in CATEGORICAL[:-2] + ["presentacion"]:
        counts = raw[column].map(repr).value_counts()
        lines.append(f"**{column}** ({raw[column].nunique()} valores): "
                     + ", ".join(f"{k} ({v})" for k, v in counts.items()))
    lines.append(f"**punto_venta_id**: {raw['punto_venta_id'].nunique()} valores")
    for column in ["unidades", "venta_neta_cop"]:
        values = _numeric(raw[column])
        lines.append(f"**{column}**: no numéricos {values.isna().sum()}, min {values.min():,.0f}, "
                     f"mediana {values.median():,.0f}, max {values.max():,.0f}")
    return "\n\n".join(lines)


def rules_report(raw: pd.DataFrame, brand: str, periods: dict[str, list[str]]) -> str:
    canonical = _canonical_spelling(raw["marca"])
    is_brand = raw["marca"].map(canonical).eq(brand)
    period = assign_period(raw.assign(mes=raw["mes"].str.strip()), periods)
    sales = _numeric(raw["venta_neta_cop"])
    rows = []
    blocks = []
    for rule in build_rules(raw):
        mask = rule.detect(raw)
        hits = raw[mask]
        brand_sales = sales[mask & is_brand].fillna(0).sum()
        not_quantifiable = (mask & is_brand & sales.isna()).any()
        rows.append({"código": rule.code, "regla": rule.name, "filas": int(mask.sum()),
                     f"filas {brand}": int((mask & is_brand).sum()),
                     f"venta {brand} afectada (COP)": "no cuantificable" if not_quantifiable else f"{brand_sales:,.0f}",
                     "propuesta": rule.proposal})
        if len(hits):
            sample = hits.assign(periodo=period[mask]).head(5)
            blocks.append(f"#### {rule.code} · {rule.name}\n\n{_fmt(sample)}")
    return _md(pd.DataFrame(rows).set_index("código")) + "\n\n" + "\n\n".join(blocks)


def variation_by(df: pd.DataFrame, brand: str, periods: dict[str, list[str]],
                 dims: Sequence[str], value: str = "v") -> pd.DataFrame:
    first, second = list(periods)
    data = df[df["marca"] == brand].assign(periodo=assign_period(df[df["marca"] == brand], periods))
    table = data.pivot_table(index=list(dims), columns="periodo", values=value, aggfunc="sum")
    table["var_abs"] = table[second] - table[first]
    table["var_pct"] = (table[second] / table[first] - 1) * 100
    return table.sort_values("var_pct")


def sensitivity(raw: pd.DataFrame, brand: str, periods: dict[str, list[str]]) -> str:
    """Variación % por región con y sin cada paso de limpieza, para ver si cambia el ranking."""
    variants: dict[str, pd.DataFrame] = {
        "todas las reglas": clean(raw),
        "sin eliminar duplicados": clean(raw, dedup=False),
        "sin normalizar dimensiones": clean(raw, normalize=False),
        "imputando venta faltante (escenario)": clean(raw, impute_missing_sales=True),
    }
    columns = {name: variation_by(frame, brand, periods, ["region"])["var_pct"]
               for name, frame in variants.items()}
    # Sin normalizar, las variantes con espacios quedan como filas aparte (se muestran entre comillas).
    table = pd.DataFrame(columns).round(2).sort_values("todas las reglas")
    table.index = [repr(r) if r != r.strip() else r for r in table.index]
    return _md(table)


def extract_documents(raw_dir: Path, output_dir: Path) -> str:
    lines = []
    for path in sorted(raw_dir.iterdir()):
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            with pymupdf.open(path) as document:
                text = "\n".join(f"--- página {i + 1} ---\n{page.get_text()}" for i, page in enumerate(document))
                detail = f"{document.page_count} páginas"
        elif suffix == ".docx":
            document = docx.Document(str(path))
            paragraphs = [f"[{p.style.name}] {p.text}" for p in document.paragraphs if p.text.strip()]
            tables = ["\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
                      for table in document.tables]
            text = "\n".join(paragraphs) + "".join(f"\n--- tabla {i + 1} ---\n{t}" for i, t in enumerate(tables))
            detail = f"{len(paragraphs)} párrafos, {len(tables)} tablas"
        elif suffix == ".md":
            text = path.read_text(encoding="utf-8")
            detail = f"{len(text.splitlines())} líneas"
        else:
            continue
        (output_dir / f"{path.name}.txt").write_text(text, encoding="utf-8")
        ids = sorted({i for i in re.findall(r"\b(?:REP|INV|OC|CAM|ENT)-[A-Z0-9-]+\b", text)
                      if not PDV_PATTERN.match(i)})
        lines.append(f"- `{path.name}`: {detail}, {len(text.split())} palabras; IDs de evidencia: "
                     + (", ".join(ids) if ids else "ninguno"))
    return "\n".join(lines)


# --- CLI ---------------------------------------------------------------------


def _parse_period(value: str) -> tuple[str, list[str]]:
    name, _, months = value.partition("=")
    parsed = [m.strip() for m in months.split(",") if m.strip()]
    if not name or not parsed or not all(MONTH_PATTERN.match(m) for m in parsed):
        raise argparse.ArgumentTypeError(f"Período inválido: {value!r}. Usa NOMBRE=AAAA-MM,AAAA-MM")
    return name, parsed


def parse_args(argv: Sequence[str] | None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--raw-dir", type=Path, required=True)
    parser.add_argument("--sales-file", required=True)
    parser.add_argument("--brand", required=True)
    parser.add_argument("--period", type=_parse_period, action="append", required=True,
                        help="Dos veces, en orden: base y comparación (p. ej. T1=... y T2=...)")
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args(argv)
    if len(args.period) != 2:
        parser.error("Indica exactamente dos --period (base y comparación).")
    return args


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    periods = dict(args.period)
    sales_path = args.raw_dir / args.sales_file
    if not sales_path.is_file():
        print(f"No existe el archivo de ventas: {sales_path}", file=sys.stderr)
        return 1
    args.output_dir.mkdir(parents=True, exist_ok=True)

    raw = load_raw(sales_path)
    cleaned = clean(raw)
    if args.brand not in set(cleaned["marca"]):
        print(f"La marca {args.brand!r} no aparece en los datos limpios.", file=sys.stderr)
        return 1
    first, second = list(periods)
    by_region = variation_by(cleaned, args.brand, periods, ["region"])
    total = variation_by(cleaned.assign(total="total"), args.brand, periods, ["total"])
    all_brands = pd.concat({b: variation_by(cleaned, b, periods, ["marca"]).iloc[0]
                            for b in sorted(cleaned["marca"].unique())}, axis=1).T

    sections = {
        "Perfil del CSV (valores crudos)": profile(raw),
        "Reglas de calidad e impacto": rules_report(raw, args.brand, periods),
        f"Variación de venta neta de {args.brand} por región ({second} vs {first})": _md(by_region.round(2)),
        f"Total {args.brand}": _md(total.round(2)),
        "Sensibilidad del ranking regional a la limpieza": sensitivity(raw, args.brand, periods),
        f"{args.brand} por canal": _md(variation_by(cleaned, args.brand, periods, ["canal"]).round(2)),
        f"{args.brand} por presentación": _md(variation_by(cleaned, args.brand, periods, ["presentacion"]).round(2)),
        f"{args.brand} por región y canal": _md(variation_by(cleaned, args.brand, periods, ["region", "canal"]).round(2)),
        f"{args.brand} por región y presentación": _md(
            variation_by(cleaned, args.brand, periods, ["region", "presentacion"]).round(2)),
        f"{args.brand} unidades por región": _md(variation_by(cleaned, args.brand, periods, ["region"], "u").round(2)),
        f"{args.brand} por punto de venta (resumen de var_pct por región)": _md(
            variation_by(cleaned, args.brand, periods, ["region", "punto_venta_id"])
            .groupby(level="region")["var_pct"].describe().round(2)),
        f"{args.brand} venta neta mensual por región": _md(
            cleaned[cleaned["marca"] == args.brand].pivot_table(index="region", columns="mes", values="v", aggfunc="sum")),
        "Todas las marcas (total)": _md(all_brands[[first, second, "var_abs", "var_pct"]].round(2)),
        "Documentos": extract_documents(args.raw_dir, args.output_dir),
    }
    report = f"# Reporte de exploración\n\nMarca: {args.brand} · Períodos: {periods}\n\n" + "\n\n".join(
        f"## {title}\n\n{body}" for title, body in sections.items())
    (args.output_dir / "report.md").write_text(report, encoding="utf-8")
    print(report)
    return 0


if __name__ == "__main__":
    sys.exit(main())
