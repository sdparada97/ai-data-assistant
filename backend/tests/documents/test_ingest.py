"""Requisitos: Publicación atómica e idempotente, Trazabilidad de la carga de documentos,
Embeddings de los fragmentos (guardados) y errores que no publican la base."""

from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import docx
import duckdb
import pymupdf
import pytest

from analyst_backend.documents import ingest as ingest_module
from analyst_backend.documents.errors import ConversionError, DocumentsError, EmbedderError, SourceError
from analyst_backend.documents.ingest import DocumentsSummary, ingest_documents

from .conftest import RAW_DIR, documents_settings, fake_embedder


def _rows(db: Path, query: str) -> list[tuple]:
    with duckdb.connect(str(db), read_only=True) as connection:
        return connection.execute(query).fetchall()


def _copy_raw(target: Path) -> Path:
    target.mkdir()
    for path in RAW_DIR.iterdir():
        if path.suffix.lower() in {".pdf", ".docx", ".md"}:
            shutil.copy(path, target / path.name)
    return target


# --- Embeddings guardados y Markdown guardado --------------------------------------------------


def test_every_chunk_has_a_384_value_embedding(fake_db: DocumentsSummary) -> None:
    rows = _rows(fake_db.documents_db_path, "SELECT len(embedding) FROM doc_chunks")

    assert len(rows) == 55
    assert {length for (length,) in rows} == {384}


def test_markdown_is_stored(fake_db: DocumentsSummary) -> None:
    stored = dict(_rows(fake_db.documents_db_path, "SELECT doc_id, markdown FROM documents"))

    assert len(stored) == 4
    assert all(markdown.strip() for markdown in stored.values())
    assert stored["01_Contexto_Negocio"] == (RAW_DIR / "01_Contexto_Negocio.md").read_text(encoding="utf-8")


# --- Publicación atómica e idempotente ---------------------------------------------------------


def test_double_load_is_identical(tmp_path: Path, fake_config_path: Path, fake_db: DocumentsSummary) -> None:
    again = ingest_documents(documents_settings(tmp_path / "otra.duckdb", fake_config_path), fake_embedder())

    for table in ("documents", "doc_chunks"):
        query = f"SELECT * FROM {table} ORDER BY 1"
        assert _rows(again.documents_db_path, query) == _rows(fake_db.documents_db_path, query)
    assert again.ingestion_id == fake_db.ingestion_id


def test_failure_mid_load_keeps_previous_database(
    tmp_path: Path, fake_config_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    db = tmp_path / "documents.duckdb"
    settings = documents_settings(db, fake_config_path)
    ingest_documents(settings, fake_embedder())
    before = _rows(db, "SELECT chunk_id, text FROM doc_chunks ORDER BY 1")
    real_write = ingest_module.write_tables

    def write_then_fail(*args: object) -> None:
        real_write(*args)
        raise RuntimeError("fallo forzado después de escribir")

    monkeypatch.setattr(ingest_module, "write_tables", write_then_fail)
    with pytest.raises(RuntimeError, match="fallo forzado"):
        ingest_documents(settings, fake_embedder())

    assert _rows(db, "SELECT chunk_id, text FROM doc_chunks ORDER BY 1") == before
    assert sorted(p.name for p in tmp_path.iterdir()) == ["documents.duckdb"]


# --- Errores que no publican la base -----------------------------------------------------------


def _empty_pdf(path: Path) -> None:
    with pymupdf.open() as document:
        page = document.new_page()
        page.draw_rect(pymupdf.Rect(50, 50, 200, 100), fill=(0.9, 0.9, 0.9))
        document.save(path)


@pytest.mark.parametrize(
    ("files", "error", "fragment"),
    [
        ({"informe.pdf": None, "informe.md": "# x\n\n## A\n\ntexto\n"}, SourceError, "informe"),
        ({"ventas.csv": "a,b\n"}, SourceError, "No hay documentos"),
        ({"roto.pdf": "no es un PDF"}, ConversionError, "roto.pdf"),
        ({"roto.docx": "no es un Word"}, ConversionError, "roto.docx"),
        ({"titulos.md": "# Título\n\n## A\n"}, ConversionError, "titulos"),
        ({"vacio.pdf": None}, ConversionError, "vacio"),
    ],
    ids=["doc_id-duplicado", "sin-documentos", "pdf-danado", "docx-danado", "solo-titulos", "pdf-sin-texto"],
)
def test_input_errors_do_not_publish(
    tmp_path: Path, fake_config_path: Path, files: dict[str, str | None], error: type[Exception], fragment: str
) -> None:
    source = tmp_path / "docs"
    source.mkdir()
    for name, content in files.items():
        if content is None:  # PDF válido sin texto (solo un rectángulo)
            _empty_pdf(source / name)
        else:
            (source / name).write_text(content, encoding="utf-8")
    db = tmp_path / "documents.duckdb"

    with pytest.raises(error, match=fragment):
        ingest_documents(documents_settings(db, fake_config_path, source), fake_embedder())

    assert not db.exists()


def test_missing_model_does_not_create_the_database(tmp_path: Path, config_path_real: Path) -> None:
    db = tmp_path / "documents.duckdb"

    with pytest.raises(EmbedderError) as error:
        ingest_documents(documents_settings(db, config_path_real, models_dir=tmp_path / "vacia"))

    assert "make models" in str(error.value) and "MODELS_CACHE_DIR" in str(error.value)
    assert not db.exists()


def test_missing_model_does_not_modify_an_existing_database(
    tmp_path: Path, fake_config_path: Path, config_path_real: Path
) -> None:
    db = tmp_path / "documents.duckdb"
    ingest_documents(documents_settings(db, fake_config_path), fake_embedder())
    before = db.read_bytes()

    with pytest.raises(EmbedderError):
        ingest_documents(documents_settings(db, config_path_real, models_dir=tmp_path / "vacia"))

    assert db.read_bytes() == before


def test_file_errors_come_before_chunking_errors(tmp_path: Path, fake_config_path: Path) -> None:
    # "a_titulos" va antes en orden alfabético, pero un archivo dañado se detecta antes que un documento vacío.
    source = tmp_path / "docs"
    source.mkdir()
    (source / "a_titulos.md").write_text("# Título\n\n## A\n", encoding="utf-8")
    (source / "b_roto.pdf").write_text("no es un PDF", encoding="utf-8")

    with pytest.raises(ConversionError, match="b_roto.pdf"):
        ingest_documents(documents_settings(tmp_path / "d.duckdb", fake_config_path, source), fake_embedder())


def test_errors_keep_an_existing_database(tmp_path: Path, fake_config_path: Path) -> None:
    db = tmp_path / "documents.duckdb"
    ingest_documents(documents_settings(db, fake_config_path), fake_embedder())
    before = db.read_bytes()
    broken = tmp_path / "docs"
    broken.mkdir()
    (broken / "roto.pdf").write_text("no es un PDF", encoding="utf-8")

    with pytest.raises(DocumentsError):
        ingest_documents(documents_settings(db, fake_config_path, broken), fake_embedder())

    assert db.read_bytes() == before


# --- Trazabilidad de la carga de documentos ----------------------------------------------------


def test_load_is_recorded(fake_db: DocumentsSummary, fake_config_path: Path) -> None:
    [(ingestion_id, model, model_sha, dimension, total, files)] = _rows(
        fake_db.documents_db_path,
        "SELECT ingestion_id, model_name, model_sha256, embedding_dimension, chunk_count, files FROM doc_ingestion_runs",
    )

    assert ingestion_id == fake_db.ingestion_id
    assert (model, model_sha, dimension, total) == ("pruebas/fake-embedder", "f" * 64, 384, 55)
    assert [f["file_name"] for f in files] == [
        "01_Contexto_Negocio.md",
        "03_Investigacion_Consumidor.pdf",
        "04_Observaciones_Canales.docx",
        "05_Reporte_Comercial.pdf",
    ]
    for record in files:
        assert record["sha256"] == hashlib.sha256((RAW_DIR / record["file_name"]).read_bytes()).hexdigest()
    assert sum(f["chunk_count"] for f in files) == 55


def test_renamed_file_changes_ingestion_id(tmp_path: Path, fake_config_path: Path, fake_db: DocumentsSummary) -> None:
    source = _copy_raw(tmp_path / "docs")
    (source / "01_Contexto_Negocio.md").rename(source / "01_Contexto.md")

    renamed = ingest_documents(documents_settings(tmp_path / "r.duckdb", fake_config_path, source), fake_embedder())

    assert renamed.ingestion_id != fake_db.ingestion_id


def test_word_document_is_loaded_from_a_copy(tmp_path: Path, fake_config_path: Path) -> None:
    # Un Word mínimo creado en la prueba recorre todo el flujo.
    source = tmp_path / "docs"
    source.mkdir()
    document = docx.Document()
    document.add_heading("Nota", level=1)
    document.add_paragraph("REP-01 | Afirmación de prueba.")
    document.save(str(source / "nota.docx"))

    summary = ingest_documents(documents_settings(tmp_path / "d.duckdb", fake_config_path, source), fake_embedder())

    assert summary.chunks_per_document == {"nota": 1}
