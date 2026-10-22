"""Requisitos: Conversión a Markdown, Estructura de los PDF y Tablas de PDF."""

from __future__ import annotations

from pathlib import Path

import docx
import pymupdf
import pytest

from analyst_backend.documents.config import DocumentsConfig
from analyst_backend.documents.errors import ConversionError
from analyst_backend.documents.markdown import docx_to_markdown, pdf_to_markdown, to_markdown
from analyst_backend.documents.sources import SourceFile

from .conftest import RAW_DIR


def _pdf(config: DocumentsConfig, name: str) -> str:
    return pdf_to_markdown(RAW_DIR / name, config.pdf, config.id_start_regex)


def _write_pdf(path: Path, pages: list[list[str]]) -> Path:
    with pymupdf.open() as document:
        for lines in pages:
            page = document.new_page()
            for index, text in enumerate(lines):
                page.insert_text((72, 72 + 40 * index), text, fontsize=10)
        document.save(path)
    return path


# --- Conversión a Markdown ---------------------------------------------------------------------


def test_word_heading_levels() -> None:
    markdown = docx_to_markdown(RAW_DIR / "04_Observaciones_Canales.docx")

    lines = markdown.splitlines()
    assert "# Observaciones de canales" in lines
    assert "## Registro de hallazgos por región" in lines
    assert "### OC-SUR-01 | Sur" in lines


def test_word_merged_cell_appears_once(tmp_path: Path) -> None:
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Concepto", "Valor"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Total"
    path = tmp_path / "combinada.docx"
    document.save(str(path))

    markdown = docx_to_markdown(path)

    assert "| Total |  |" in markdown.splitlines()
    assert markdown.count("Total") == 1


def test_repeated_spaces_are_collapsed() -> None:
    markdown = docx_to_markdown(RAW_DIR / "04_Observaciones_Canales.docx")

    row = next(line for line in markdown.splitlines() if "CAM-101" in line)
    assert "CAM-101 2026-04-09" in row
    assert "  " not in row.replace("|  |", "")


def test_markdown_file_is_returned_as_is(config: DocumentsConfig) -> None:
    path = RAW_DIR / "01_Contexto_Negocio.md"
    source = SourceFile(doc_id=path.stem, path=path, source_type="md")

    assert to_markdown(source, config) == path.read_text(encoding="utf-8")


@pytest.mark.parametrize(("name", "content"), [("roto.pdf", b"esto no es un PDF"), ("roto.docx", b"tampoco un Word")])
def test_damaged_file_names_the_file(tmp_path: Path, config: DocumentsConfig, name: str, content: bytes) -> None:
    path = tmp_path / name
    path.write_bytes(content)
    source = SourceFile(doc_id=path.stem, path=path, source_type=path.suffix.lstrip("."))

    with pytest.raises(ConversionError, match=name):
        to_markdown(source, config)


# --- Estructura de los PDF ---------------------------------------------------------------------


def test_pdf_headers_and_footers_are_removed(config: DocumentsConfig) -> None:
    markdown = _pdf(config, "05_Reporte_Comercial.pdf")

    assert "ALMENA FOODS / ANÁLISIS COMERCIAL" not in markdown
    assert "Página 1" not in markdown
    assert "<!-- page: 2 -->" in markdown.splitlines()
    assert "## 3. Interpretación preliminar" in markdown.splitlines()
    assert markdown.splitlines()[2] == "# Reporte ejecutivo | Desempeño de Brío"


def test_pdf_paragraph_lines_are_joined(config: DocumentsConfig) -> None:
    markdown = _pdf(config, "05_Reporte_Comercial.pdf")

    assert (
        "REP-01 | Localización de la presión comercial. El reporte considera a Centro la región con mayor caída "
        "porcentual y recomienda intervenir allí primero."
    ) in markdown.splitlines()


def test_single_page_pdf_keeps_its_text(tmp_path: Path, config: DocumentsConfig) -> None:
    path = _write_pdf(tmp_path / "una.pdf", [["Línea A", "Línea B"]])

    markdown = pdf_to_markdown(path, config.pdf, config.id_start_regex)

    assert "Línea A" in markdown
    assert "Línea B" in markdown


def test_line_repeated_on_a_single_page_is_kept(tmp_path: Path, config: DocumentsConfig) -> None:
    path = _write_pdf(tmp_path / "dos.pdf", [["Repetida", "Otra", "Repetida"], ["Segunda página"]])

    markdown = pdf_to_markdown(path, config.pdf, config.id_start_regex)

    assert markdown.count("Repetida") == 2


# --- Tablas de PDF -----------------------------------------------------------------------------


def test_pdf_table_without_borders(config: DocumentsConfig) -> None:
    lines = _pdf(config, "05_Reporte_Comercial.pdf").splitlines()

    assert "| Región | Variación reportada (T2 vs. T1) | Lectura del analista |" in lines
    assert "| Centro | -23 % | Mayor caída porcentual; prioridad principal. |" in lines


def test_pdf_cell_with_several_lines(config: DocumentsConfig) -> None:
    markdown = _pdf(config, "03_Investigacion_Consumidor.pdf")

    section = markdown.split("## 3. Hallazgos por región", 1)[1].split("## 4.", 1)[0]
    assert (
        "| Centro | Algunos comentarios sobre promociones ajenas y faltantes ocasionales. "
        "| No respalda inferir la mayor caída en ventas. |"
    ) in section.splitlines()


# --- Coherencia entre configuración y conversores ----------------------------------------------


def test_every_supported_extension_has_a_converter(tmp_path: Path, config: DocumentsConfig) -> None:
    from analyst_backend.documents.config import SUPPORTED_EXTENSIONS

    samples = {".pdf": _write_pdf(tmp_path / "a.pdf", [["Texto"]]), ".md": tmp_path / "a.md", ".docx": tmp_path / "a.docx"}
    samples[".md"].write_text("# Título\n\n## S\n\nTexto\n", encoding="utf-8")
    document = docx.Document()
    document.add_paragraph("Texto")
    document.save(str(samples[".docx"]))

    assert set(samples) == set(SUPPORTED_EXTENSIONS)
    for extension, path in samples.items():
        source = SourceFile(doc_id=path.stem, path=path, source_type=extension.lstrip("."))
        assert "Texto" in to_markdown(source, config)
